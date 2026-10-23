"""Adaptadores próprios e locais para documentos OOXML e PDF.

Nenhum script dos repositórios de referência é importado ou executado.
"""
from __future__ import annotations

import csv
import io
import re
import tempfile
import zipfile
from itertools import islice
from pathlib import Path
from typing import Any, Callable
from xml.sax.saxutils import escape

from .workspace import DOCUMENT_EXTENSIONS, MAX_DOCUMENT_BYTES, MAX_DOCUMENT_TEXT, download_url, regular_file, safe_path


def _archive_guard(path: Path) -> None:
    """Limita expansão de OOXML antes de entregar o arquivo às bibliotecas."""
    if not zipfile.is_zipfile(path):
        raise ValueError("O documento não possui uma estrutura OOXML válida.")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 2_048 or sum(entry.file_size for entry in entries) > 100_000_000:
                raise ValueError("O documento expandido excede os limites de segurança.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError("Documentos OOXML criptografados não são suportados.")
    except zipfile.BadZipFile as exc:
        raise ValueError("O documento OOXML está corrompido.") from exc


def _input_file(root: Path, relative: str, extension: str) -> Path:
    path = regular_file(root, relative, {extension}, MAX_DOCUMENT_BYTES)
    if extension in {".docx", ".xlsx", ".pptx"}:
        _archive_guard(path)
    return path


def _output_file(root: Path, relative: str, extension: str, overwrite: bool) -> Path:
    path = safe_path(root, relative)
    if path.suffix.lower() != extension:
        raise ValueError(f"O destino precisa usar a extensão {extension}.")
    if path.exists() and (not path.is_file() or path.is_symlink()):
        raise ValueError("O destino existe, mas não é um arquivo regular.")
    if path.exists() and not overwrite:
        raise ValueError("O arquivo já existe. Ative 'Permitir sobrescrever' para substituí-lo.")
    return path


def _atomic_write(path: Path, writer: Callable[[Path], None]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_name = ""
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix=path.suffix, delete=False) as temporary:
            temporary_name = temporary.name
        writer(Path(temporary_name))
        temporary_path = Path(temporary_name)
        size = temporary_path.stat().st_size
        if not 20 < size <= MAX_DOCUMENT_BYTES:
            raise ValueError("O documento gerado excede o limite de tamanho.")
        temporary_path.replace(path)
        return size
    finally:
        if temporary_name:
            Path(temporary_name).unlink(missing_ok=True)


def _limited(text: str) -> tuple[str, bool]:
    truncated = len(text) > MAX_DOCUMENT_TEXT
    return text[:MAX_DOCUMENT_TEXT] + ("\n\n[Extração limitada a 100.000 caracteres.]" if truncated else ""), truncated


def read_document(root: Path, relative: str) -> dict[str, Any]:
    extension = safe_path(root, relative).suffix.lower()
    if extension not in DOCUMENT_EXTENSIONS:
        raise ValueError("Use um arquivo DOCX, XLSX, PPTX ou PDF.")
    path = _input_file(root, relative, extension)
    metadata: dict[str, Any] = {}
    if extension == ".docx":
        from docx import Document

        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables[:50]:
            for row in table.rows[:500]:
                parts.append(" | ".join(cell.text for cell in row.cells))
        metadata = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        parts = []
        rows_read = 0
        try:
            for worksheet in workbook.worksheets[:20]:
                parts.append(f"## Planilha: {worksheet.title}")
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    if row_index > 500 or rows_read >= 10_000:
                        break
                    values = ["" if value is None else str(value) for value in row[:100]]
                    if any(values):
                        parts.append(" | ".join(values).rstrip(" |"))
                        rows_read += 1
        finally:
            workbook.close()
        metadata = {"sheets": len(workbook.worksheets), "rowsRead": rows_read}
    elif extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        parts = []
        for index, slide in enumerate(islice(presentation.slides, 100), start=1):
            parts.append(f"## Slide {index}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        metadata = {"slides": len(presentation.slides)}
    else:
        from pypdf import PdfReader

        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            raise ValueError("PDF criptografado não é suportado.")
        parts = []
        for index, page in enumerate(islice(reader.pages, 200), start=1):
            parts.append(f"## Página {index}\n{page.extract_text() or ''}")
        metadata = {"pages": len(reader.pages)}
    text, truncated = _limited("\n\n".join(parts).strip())
    metadata["truncated"] = truncated
    return {
        "path": path.relative_to(root.resolve()).as_posix(),
        "format": extension[1:],
        "bytes": path.stat().st_size,
        "text": text or "[Nenhum texto extraível encontrado.]",
        "downloadUrl": download_url(path.relative_to(root.resolve()).as_posix()),
        **metadata,
    }


def _docx_writer(path: Path, title: str, content: str) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.core_properties.title = title
    document.core_properties.author = "JARVIS"
    document.add_heading(title, 0)
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(stripped)
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Gerado localmente pelo JARVIS"
        footer.runs[0].font.size = Pt(8)
    document.save(path)


def _xlsx_rows(content: str) -> list[list[str]]:
    first_line = next((line for line in content.splitlines() if line.strip()), "")
    delimiter = max((",", ";", "\t"), key=lambda candidate: first_line.count(candidate))
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(content), delimiter=delimiter) if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError("Informe ao menos uma linha de dados.")
    if len(rows) > 1_000 or any(len(row) > 100 for row in rows):
        raise ValueError("A planilha excede o limite de 1.000 linhas ou 100 colunas.")
    return rows


def _xlsx_writer(path: Path, title: str, content: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _xlsx_rows(content)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = re.sub(r"[\\/*?:\[\]]", "-", title).strip()[:31] or "Dados"
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value[:32_767])
            # Entradas são dados, nunca fórmulas ou comandos de planilha.
            cell.data_type = "s"
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24352B")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index in range(1, max(len(row) for row in rows) + 1):
        width = min(40, max(10, max(len(row[column_index - 1]) if column_index <= len(row) else 0 for row in rows) + 2))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    workbook.save(path)


def _pptx_writer(path: Path, title: str, content: str) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    sections = [section for section in re.split(r"(?m)^\s*---\s*$", content) if section.strip()]
    if not sections:
        raise ValueError("Informe o conteúdo dos slides.")
    if len(sections) > 25:
        raise ValueError("Use no máximo 25 slides por apresentação.")
    presentation = Presentation()
    presentation.core_properties.title = title
    presentation.core_properties.author = "JARVIS"
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, section in enumerate(sections, start=1):
        lines = [line.strip() for line in section.splitlines() if line.strip()]
        slide_title = lines[0].lstrip("# ") if lines else title
        bullets = [line[2:] if line.startswith("- ") else line for line in lines[1:21]]
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(15, 22, 18)
        label_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.25), Inches(11.8), Inches(0.3))
        label_box.text_frame.text = title
        label_box.text_frame.paragraphs[0].font.size = Pt(9)
        label_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(127, 139, 132)
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.65), Inches(11.8), Inches(1.15))
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = slide_title
        title_paragraph.font.size = Pt(30)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = RGBColor(185, 245, 122)
        body_box = slide.shapes.add_textbox(Inches(0.85), Inches(2.0), Inches(11.4), Inches(4.4))
        body = body_box.text_frame
        body.word_wrap = True
        body.text = bullets[0] if bullets else ""
        for bullet in (bullets[1:] if bullets else []):
            body.add_paragraph().text = bullet
        for paragraph in body.paragraphs:
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = RGBColor(237, 242, 239)
            paragraph.space_after = Pt(12)
        footer = slide.shapes.add_textbox(Inches(11.5), Inches(6.85), Inches(1.1), Inches(0.3))
        footer.text_frame.text = f"{index}/{len(sections)} · JARVIS"
        footer.text_frame.paragraphs[0].font.size = Pt(9)
        footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(127, 139, 132)
    presentation.save(path)


def _pdf_writer(path: Path, title: str, content: str) -> None:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("JarvisTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=25, alignment=TA_LEFT)
    heading_style = ParagraphStyle("JarvisHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=17, spaceBefore=8, spaceAfter=5)
    body_style = ParagraphStyle("JarvisBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, spaceAfter=6)
    story = [Paragraph(escape(title), title_style), Spacer(1, 7 * mm)]
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 2 * mm))
        elif stripped.startswith("#"):
            story.append(Paragraph(escape(stripped.lstrip("# ")), heading_style))
        elif stripped.startswith("- "):
            story.append(Paragraph(escape(stripped[2:]), body_style, bulletText="•"))
        else:
            story.append(Paragraph(escape(stripped), body_style))
    document = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=title, author="JARVIS")
    document.build(story)


_WRITERS: dict[str, Callable[[Path, str, str], None]] = {
    "docx": _docx_writer,
    "xlsx": _xlsx_writer,
    "pptx": _pptx_writer,
    "pdf": _pdf_writer,
}


def create_document(root: Path, document_format: str, relative: str, title: str, content: str, overwrite: bool, dry_run: bool) -> dict[str, Any]:
    if document_format not in _WRITERS:
        raise ValueError("Formato de saída não suportado.")
    extension = f".{document_format}"
    path = _output_file(root, relative, extension, overwrite)
    normalized = path.relative_to(root.resolve()).as_posix()
    if document_format == "xlsx":
        _xlsx_rows(content)
    if dry_run:
        return {"dryRun": True, "wouldWrite": normalized, "format": document_format, "overwrite": path.exists(), "characters": len(content)}
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and not overwrite:
        raise ValueError("O arquivo já existe. Ative 'Permitir sobrescrever' para substituí-lo.")
    size = _atomic_write(path, lambda temporary: _WRITERS[document_format](temporary, title, content))
    return {"dryRun": False, "written": normalized, "format": document_format, "bytes": size, "downloadUrl": download_url(normalized)}


def merge_pdfs(root: Path, paths: str, output: str, overwrite: bool, dry_run: bool) -> dict[str, Any]:
    from pypdf import PdfReader, PdfWriter

    relatives = [line.strip() for line in paths.splitlines() if line.strip()]
    if not 2 <= len(relatives) <= 10 or len(set(relatives)) != len(relatives):
        raise ValueError("Informe de 2 a 10 PDFs distintos, um por linha.")
    inputs = [_input_file(root, relative, ".pdf") for relative in relatives]
    readers = [PdfReader(path, strict=False) for path in inputs]
    if any(reader.is_encrypted for reader in readers):
        raise ValueError("PDFs criptografados não são suportados.")
    pages = sum(len(reader.pages) for reader in readers)
    if pages > 500:
        raise ValueError("A mesclagem excede o limite de 500 páginas.")
    destination = _output_file(root, output, ".pdf", overwrite)
    normalized = destination.relative_to(root.resolve()).as_posix()
    if destination in inputs:
        raise ValueError("O destino não pode ser um dos PDFs de entrada.")
    if dry_run:
        return {"dryRun": True, "wouldWrite": normalized, "sources": relatives, "pages": pages, "overwrite": destination.exists()}
    size = _atomic_write(destination, lambda temporary: _write_merged(readers, temporary))
    return {"dryRun": False, "written": normalized, "format": "pdf", "bytes": size, "pages": pages, "downloadUrl": download_url(normalized)}


def _write_merged(readers: list[Any], path: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for reader in readers:
        for page in reader.pages:
            writer.add_page(page)
    with path.open("wb") as output:
        writer.write(output)
