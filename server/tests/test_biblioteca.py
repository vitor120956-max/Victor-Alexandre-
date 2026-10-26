"""Testes da biblioteca (0.51.0): aprender com livro, vídeo e pesquisa.

Sem rede e sem modelo: os arquivos são montados na hora (PDF com reportlab,
EPUB e DOCX com as bibliotecas já instaladas) e a busca é BM25 puro.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from server import biblioteca
from server.biblioteca import (
    BibliotecaError,
    buscar,
    correlacoes,
    digest,
    gravar_digest,
    importar,
    importar_texto,
    listar,
    obter,
    remover,
    responder,
)

TEXTO_HTML = """# Capítulo 1 — Fundamentos

A contabilidade de custo separa gasto de despesa e mostra o que consome margem.
Cada lançamento precisa de documento de apoio, data e responsável.

# Capítulo 2 — Rateio

O rateio distribui o custo indireto entre os produtos. O critério precisa ser
defensável, estável e documentado; mudar o critério todo mês esconde o problema.
"""

TEXTO_VIDEO = """WEBVTT

1
00:00:01.000 --> 00:00:04.000
Neste vídeo a gente vê como o rateio de custo indireto funciona na prática.

2
00:00:04.500 --> 00:00:08.000
O critério de rateio precisa ser estável, senão a margem do produto mente.
"""


@pytest.fixture()
def pasta(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    raiz = tmp_path / "dados"
    raiz.mkdir()
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: raiz)
    return raiz


def test_importar_texto_cria_secoes_e_indice(pasta: Path) -> None:
    fonte = importar_texto("Rateio de custo", TEXTO_HTML)
    assert fonte["secoes"] >= 2
    assert fonte["tipo"] == "pesquisa"
    assert listar()[0]["slug"] == fonte["slug"]
    assert obter(fonte["slug"])["titulo"] == "Rateio de custo"


def test_importar_texto_curto_recusa(pasta: Path) -> None:
    with pytest.raises(BibliotecaError, match="curto demais"):
        importar_texto("x", "texto curto")


def test_titulos_viram_secoes(pasta: Path) -> None:
    fonte = importar_texto("Manual", TEXTO_HTML)
    partes = biblioteca._carregar_partes(fonte["slug"])
    titulos = [parte["titulo"] for parte in partes]
    assert any("Fundamentos" in titulo for titulo in titulos)
    assert any("Rateio" in titulo for titulo in titulos)


def test_busca_acha_a_secao_certa(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    achados = buscar("critério de rateio")
    assert achados
    assert "Rateio" in achados[0]["secao"]
    assert achados[0]["cobertura"] >= 0.5


def test_busca_sem_termo_nao_acha(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    assert buscar("astronomia quântica") == []
    assert buscar("") == []


def test_responder_devolve_trecho_e_referencia(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    resposta = responder("rateio de custo indireto")
    assert resposta is not None
    assert "Manual de custos" in resposta["resposta"]
    assert "seção" in resposta["resposta"]
    assert resposta["cobertura"] >= 0.5


def test_responder_recusa_quando_a_cobertura_e_baixa(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    assert responder("rateio astronomia fotossíntese vulcão") is None


def test_correlacoes_cruzam_fontes(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    importar_texto("Nota de vídeo", "O rateio de custo indireto aparece no vídeo. " * 6)
    achados = correlacoes()
    termos = {item["termo"]: item for item in achados}
    assert "rateio" in termos
    assert len(termos["rateio"]["fontes"]) == 2


def test_correlacoes_de_um_termo_mostram_vizinhos(pasta: Path) -> None:
    importar_texto("Manual de custos", TEXTO_HTML)
    vizinhos = [item["termo"] for item in correlacoes("rateio")]
    assert "custo" in vizinhos or "indireto" in vizinhos
    assert "rateio" not in vizinhos


def test_digest_tem_front_matter_e_sumario(pasta: Path) -> None:
    fonte = importar_texto("Manual de custos", TEXTO_HTML)
    ficha = digest(fonte["slug"])
    assert ficha["markdown"].startswith("---\nname:")
    assert "## Sumário" in ficha["markdown"]
    assert "## Termos mais frequentes" in ficha["markdown"]
    assert "sem modelo de linguagem" in ficha["markdown"]


def test_gravar_digest_cria_skill_md(pasta: Path) -> None:
    fonte = importar_texto("Manual de custos", TEXTO_HTML)
    caminho = gravar_digest(fonte["slug"])
    assert caminho.name == "SKILL.md"
    assert "Manual de custos" in caminho.read_text(encoding="utf-8")


def test_remover_tira_do_indice(pasta: Path) -> None:
    fonte = importar_texto("Manual de custos", TEXTO_HTML)
    assert remover(fonte["slug"]) is True
    assert listar() == []
    assert remover(fonte["slug"]) is False


# ------------------------------------------------------------------ formatos

def test_importar_vtt_tira_tempo_e_numero(pasta: Path, tmp_path: Path) -> None:
    arquivo = tmp_path / "aula.vtt"
    arquivo.write_text(TEXTO_VIDEO, encoding="utf-8")
    fonte = importar(arquivo)
    assert fonte["tipo"] == "video"
    texto = " ".join(parte["texto"] for parte in biblioteca._carregar_partes(fonte["slug"]))
    assert "00:00" not in texto
    assert "WEBVTT" not in texto
    assert "rateio" in texto.lower()


def test_importar_pdf_de_verdade(pasta: Path, tmp_path: Path) -> None:
    from reportlab.pdfgen import canvas

    arquivo = tmp_path / "livro.pdf"
    pdf = canvas.Canvas(str(arquivo))
    pdf.drawString(72, 760, "Capitulo 1 - Custo industrial e margem de contribuicao")
    pdf.drawString(72, 740, "O custo fixo por produto depende do rateio escolhido pela empresa.")
    pdf.showPage()
    pdf.drawString(72, 760, "Capitulo 2 - Formacao de preco")
    pdf.drawString(72, 740, "Preco de venda precisa cobrir custo, despesa e margem planejada.")
    pdf.save()

    fonte = importar(arquivo)
    assert fonte["tipo"] == "livro"
    partes = biblioteca._carregar_partes(fonte["slug"])
    assert any(parte["pagina"] == 1 for parte in partes)
    assert any("rateio" in parte["texto"].lower() for parte in partes)


def test_importar_docx(pasta: Path, tmp_path: Path) -> None:
    from docx import Document

    arquivo = tmp_path / "nota.docx"
    documento = Document()
    documento.add_heading("Capítulo único", level=1)
    documento.add_paragraph("O custo indireto precisa de critério de rateio documentado e estável.")
    documento.save(str(arquivo))

    fonte = importar(arquivo)
    assert fonte["tipo"] == "documento"
    texto = " ".join(parte["texto"] for parte in biblioteca._carregar_partes(fonte["slug"]))
    assert "rateio documentado" in texto


def test_importar_epub(pasta: Path, tmp_path: Path) -> None:
    arquivo = tmp_path / "livro.epub"
    container = (
        '<?xml version="1.0"?><container><rootfiles>'
        '<rootfile full-path="OEBPS/content.opf"/></rootfiles></container>'
    )
    opf = (
        '<package><manifest><item id="c1" href="cap1.xhtml"/></manifest>'
        '<spine><itemref idref="c1"/></spine></package>'
    )
    capitulo = (
        "<html><body><h1>Capítulo 1</h1><p>O rateio do custo indireto aparece aqui com "
        "critério estável.</p></body></html>"
    )
    with zipfile.ZipFile(arquivo, "w") as pacote:
        pacote.writestr("META-INF/container.xml", container)
        pacote.writestr("OEBPS/content.opf", opf)
        pacote.writestr("OEBPS/cap1.xhtml", capitulo)

    fonte = importar(arquivo)
    texto = " ".join(parte["texto"] for parte in biblioteca._carregar_partes(fonte["slug"]))
    assert "rateio do custo indireto" in texto
    assert "<h1>" not in texto


def test_formato_nao_suportado(pasta: Path, tmp_path: Path) -> None:
    arquivo = tmp_path / "planilha.xlsx"
    arquivo.write_bytes(b"PK\x03\x04lixo")
    with pytest.raises(BibliotecaError, match="Formato não suportado"):
        importar(arquivo)


def test_arquivo_inexistente(pasta: Path, tmp_path: Path) -> None:
    with pytest.raises(BibliotecaError, match="não encontrado"):
        importar(tmp_path / "nao-existe.txt")


def test_binario_recusado(pasta: Path, tmp_path: Path) -> None:
    arquivo = tmp_path / "dados.txt"
    arquivo.write_bytes(b"\x00\x01\x02" * 200)
    with pytest.raises(BibliotecaError, match="binário"):
        importar(arquivo)


def test_biblioteca_vazia_responde_none(pasta: Path) -> None:
    assert responder("qualquer coisa") is None
    assert correlacoes() == []


def test_obter_fonte_inexistente(pasta: Path) -> None:
    with pytest.raises(BibliotecaError, match="não encontrada"):
        obter("nao-existe")


def test_partes_corrompidas_nao_derrubam_busca(pasta: Path) -> None:
    fonte = importar_texto("Manual", TEXTO_HTML)
    biblioteca._arquivo_da_fonte(fonte["slug"]).write_text("{lixo", encoding="utf-8")
    assert buscar("rateio") == []


# -------------------------------------------------------- sumário e exportar

def test_sumario_lista_as_secoes_com_pagina_e_termos(tmp_path, monkeypatch):
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: tmp_path)
    fonte = importar_texto(
        "Manual de custo",
        "# Capítulo 1 — Rateio\n\nO rateio distribui o custo indireto entre os produtos fabricados. "
        "O critério precisa ser defensável e documentado pela contabilidade.\n\n"
        "# Capítulo 2 — Margem\n\nA margem de contribuição é a receita menos o custo variável de cada produto.",
    )
    indice = biblioteca.sumario(fonte["slug"])
    assert indice["secoes"] >= 2
    titulos = [capitulo["titulo"] for capitulo in indice["capitulos"]]
    assert any("Rateio" in titulo for titulo in titulos)
    primeiro = indice["capitulos"][0]
    assert primeiro["comeca_com"]
    assert primeiro["termos"]


def test_sumario_de_fonte_inexistente_recusa(tmp_path, monkeypatch):
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: tmp_path)
    with pytest.raises(BibliotecaError):
        biblioteca.sumario("nao-existe")


def test_exportar_grava_markdown_com_ficha_e_texto(tmp_path, monkeypatch):
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: tmp_path)
    fonte = importar_texto(
        "Nota de rateio",
        "O rateio distribui o custo indireto entre os produtos fabricados conforme o critério "
        "escolhido pela contabilidade da empresa, e o total rateado fecha com o total da conta "
        "de custos do período. O critério precisa ser defensável, estável e documentado, porque "
        "mudar de critério muda o custo de cada produto e o preço de venda.",
    )
    resultado = biblioteca.exportar(fonte["slug"])
    arquivo = Path(resultado["arquivo"])
    assert arquivo.is_file()
    texto = arquivo.read_text(encoding="utf-8")
    assert "## Sumário" in texto
    assert "Texto completo" in texto
    assert "rateio" in texto.lower()
    assert resultado["bytes"] > 0
    assert resultado["truncado"] is False


def test_exportar_sem_texto_sai_so_a_ficha(tmp_path, monkeypatch):
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: tmp_path)
    fonte = importar_texto(
        "Nota curta",
        "A margem de contribuição é a receita menos o custo variável de cada produto vendido pela "
        "empresa, e serve para cobrir o custo fixo e formar o lucro do período contábil. Quanto maior "
        "a margem, mais rápido o negócio cobre a estrutura que não depende do volume vendido.",
    )
    resultado = biblioteca.exportar(fonte["slug"], incluir_texto=False)
    texto = Path(resultado["arquivo"]).read_text(encoding="utf-8")
    assert "Texto completo" not in texto
    assert "## Sumário" in texto


def test_busca_traz_contexto_maior_e_vizinhas(tmp_path, monkeypatch):
    monkeypatch.setattr(biblioteca.dados, "raiz", lambda: tmp_path)
    importar_texto(
        "Livro de custos",
        "# Capítulo 1 — Rateio\n\n" + ("O rateio distribui o custo indireto entre os produtos. " * 40) + "\n\n"
        "# Capítulo 2 — Margem\n\nA margem de contribuição é a receita menos o custo variável de cada produto.",
    )
    achado = biblioteca.buscar("rateio", limite=1)[0]
    assert len(achado["contexto"]) >= len(achado["trecho"])
    assert achado["vizinhas"]["depois"] in {None, "Capítulo 2 — Margem"}
